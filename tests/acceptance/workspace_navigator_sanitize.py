"""验收脚本：workspace_navigator 经 executor 的 server sanitize 后正文是否存活。

直接跑（不依赖 Electron / Redis）：

    .venv\\Scripts\\python.exe -m tests.acceptance.workspace_navigator_sanitize
"""

from __future__ import annotations

import asyncio
import json
from pathlib import Path

from app.agents.skills.loader import load_skill_plugins
from app.agents.skills.registry import ToolRegistry

from _paths import REPO_ROOT
ROOT = REPO_ROOT / ".ws-acceptance"


def _write_txt(path: Path) -> None:
    path.write_text(
        "第一行：项目说明\n"
        "DB_PASSWORD=hunter2xyz\n"
        "联系邮箱 zhangsan@example.com\n"
        "最后一行的正文必须存活。\n",
        encoding="utf-8",
    )


def _write_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("答辩文档", level=1)
    doc.add_paragraph("DOCX 正文段落一：系统采用前后端分离架构。")
    doc.add_paragraph("DOCX 正文段落二：token=abcdef1234567890")
    doc.save(str(path))


def _write_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
    box.text_frame.text = "PPTX 幻灯片正文：三层架构与调用流程。"
    prs.save(str(path))


def _fake_electron(handler):
    async def fake_list_tools(_server):
        return [
            {"name": "workspace_list"},
            {"name": "workspace_stat"},
            {"name": "workspace_content_extract"},
            {"name": "workspace_read"},
        ]

    async def fake_call_tool(_server, tool_name, args, **_kwargs):
        return handler(tool_name, args or {})

    return fake_list_tools, fake_call_tool


def _extract_response(path: str) -> dict:
    """模拟 Electron 的 workspace_content_extract：真实解析文件。"""
    suffix = Path(path).suffix.casefold()
    full = ROOT / path
    if suffix == ".docx":
        from docx import Document

        doc = Document(str(full))
        return {
            "status": "success",
            "data": {"paragraphs": [{"text": p.text} for p in doc.paragraphs if p.text.strip()]},
            "meta": {"format": "docx", "workspace_version": 7},
        }
    if suffix == ".pptx":
        from pptx import Presentation

        prs = Presentation(str(full))
        slides = []
        for index, slide in enumerate(prs.slides, 1):
            texts = [
                shape.text_frame.text
                for shape in slide.shapes
                if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip()
            ]
            if texts:
                slides.append({"title": f"第{index}页", "text": "\n".join(texts)})
        return {
            "status": "success",
            "data": {"slides": slides},
            "meta": {"format": "pptx", "workspace_version": 7},
        }
    return {
        "status": "success",
        "data": {"text": full.read_text(encoding="utf-8")},
        "meta": {"format": "text", "workspace_version": 7},
    }


async def _read_through_executor(path: str) -> dict:
    from app.agents.skills import executor as ex
    from app.agents.skills.capability import ToolCapability

    async def fake_capability(name, scene, user_role="user", user_id="", **_kwargs):
        return ToolCapability(
            name=name, version="1.0.0", status="stable", description="",
            category="workspace", domain="workspace",
            parameters={"type": "object", "properties": {}},
            source="mcp", environment="server", server="lumi_skill",
            raw_name="workspace_navigator", permission="user",
            write_op=False, requires_confirmation=False,
            confirmation_mode="client", idempotent=True,
            annotations={"provider": "builtin"},
        )

    original = ex.get_tool_capability
    ex.get_tool_capability = fake_capability
    try:
        result = await ex.execute_tool_call(
            {"id": "c1", "type": "function", "function": {
                "name": "workspace_navigator",
                "arguments": {"action": "read", "path": path},
            }},
            "u1", "office", "conv-1",
            authorized_workspace_id="ws-acceptance",
            allow_internal=True,
        )
    finally:
        ex.get_tool_capability = original
    return result


def _install_fake_electron(monkeypatch_like: dict, monkeypatch=None) -> None:
    """把假 Electron 接到应用模块上。

    ``monkeypatch`` 可选但**强烈建议传**：不传时是直接赋值，补丁会永久留在应用
    模块上（曾导致 ``test_workspace_read_window`` 在批次里被污染 —— 后续用例的
    ``WorkspaceReader`` 读到了假 Electron 的内容）。传了就用 pytest 的
    ``monkeypatch.setattr``，用例结束自动还原。
    """
    calls = monkeypatch_like["calls"]

    def handler(tool_name: str, args: dict) -> dict:
        calls.append({"tool": tool_name, "args": dict(args)})
        if tool_name == "workspace_list":
            entries = [
                {"path": item.name, "type": "file", "size": item.stat().st_size}
                for item in sorted(ROOT.iterdir()) if item.is_file()
            ]
            return {"status": "success", "data": {"entries": entries}, "has_more": False, "cursor": ""}
        if tool_name == "workspace_stat":
            target = ROOT / str(args.get("path") or "")
            return {"status": "success", "data": {"kind": "file", "size": target.stat().st_size}}
        return _extract_response(str(args.get("path") or ""))

    list_tools, call_tool = _fake_electron(handler)
    wn_route = {"status_code": "WORKSPACE_READY", "server_name": "lumi_pc", "device_id": "dev-1"}
    route_fn = lambda user_id, workspace_id: dict(wn_route)  # noqa: E731
    reader_route = lambda self: dict(wn_route)  # noqa: E731

    # 让 WorkspaceReader（navigator 内部解析器）走同一套假 Electron。
    import app.agents.mcp.manager as manager
    import app.workspace.context as wc
    import app.workspace.read.reader as wr

    targets = (
        (manager, "list_tools", list_tools),
        (manager, "call_tool", call_tool),
        (wc, "resolve_workspace_desktop", route_fn),
        (wr.WorkspaceReader, "_route", reader_route),
    )
    if monkeypatch is not None:
        for module, name, value in targets:
            monkeypatch.setattr(module, name, value)
        return
    for module, name, value in targets:
        setattr(module, name, value)


def collect_report(monkeypatch=None) -> dict:
    """跑完三类文件，返回可直接断言的验收报告。

    ``monkeypatch`` 由 pytest 用例传入时会自动还原假 Electron；直接当脚本跑时
    传 None（进程随即退出，不需要还原）。
    """
    load_skill_plugins()
    assert ToolRegistry.get("workspace_navigator") is not None, "workspace_navigator 未注册"

    ROOT.mkdir(parents=True, exist_ok=True)
    _write_txt(ROOT / "notes.txt")
    _write_docx(ROOT / "答辩.docx")
    _write_pptx(ROOT / "答辩.pptx")

    report: dict = {"cases": []}
    calls: list[dict] = []
    _install_fake_electron({"calls": calls}, monkeypatch)

    for path in ("notes.txt", "答辩.docx", "答辩.pptx"):
        result = asyncio.run(_read_through_executor(path))
        envelope = result.data if isinstance(result.data, dict) else {}
        data = envelope.get("data") if isinstance(envelope.get("data"), dict) else {}
        sections = data.get("sections") if isinstance(data.get("sections"), list) else []
        text = "\n".join(str(item.get("text") or "") for item in sections if isinstance(item, dict))
        report["cases"].append({
            "path": path,
            "status": envelope.get("status"),
            "executor_status": result.status,
            "path_preserved": str(data.get("path") or "") == path,
            "sections": len(sections),
            "char_count": len(text),
            "content_alive": len(text.strip()) > 0,
            "redacted": bool((envelope.get("meta") or {}).get("redacted")),
            "secret_leaked": "hunter2xyz" in text or "zhangsan@example.com" in text,
            "sample": text[:160].replace("\n", " | "),
        })
    report["electron_calls"] = calls
    return report


def main() -> int:
    report = collect_report()
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
